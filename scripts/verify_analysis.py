"""Verify source preservation and reproduce the repaired ANOVA statistics."""
from collections import defaultdict
from pathlib import Path
import json, math, statistics
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]

def verify():
    original=load_workbook(ROOT/'Bank+Churn+Data (version 1).xlsx',data_only=False)
    repaired=load_workbook(ROOT/'Bank_Churn_Analysis_Repaired.xlsx',data_only=False)
    before=list(original['Sheet1'].iter_rows(min_row=1,max_col=13,values_only=True))
    after=list(repaired['Sheet1'].iter_rows(min_row=1,max_col=13,values_only=True))
    if before!=after:raise ValueError('The original analysis input was changed.')
    rows=after[1:];groups=defaultdict(list)
    if len({r[0] for r in rows})!=len(rows):raise ValueError('Duplicate customer IDs.')
    for r in rows:
        if any(v is None for v in r):raise ValueError('Missing analysis inputs.')
        if not math.isfinite(float(r[2])):raise ValueError('Invalid credit score.')
        groups[r[3]].append(float(r[2]))
    count=len(rows);grand=statistics.mean([r[2] for r in rows])
    between=sum(len(v)*(statistics.mean(v)-grand)**2 for v in groups.values())
    within=sum(sum((x-statistics.mean(v))**2 for x in v) for v in groups.values())
    if len(groups)!=3:raise ValueError('The three-group F-tail formula requires exactly three groups.')
    df1=2;df2=count-3;f=(between/df1)/(within/df2)
    p=(df2/(df2+2*f))**(df2/2)
    if not math.isclose(repaired['RQ2']['E13'].value,f,rel_tol=1e-10):raise ValueError('ANOVA F value differs.')
    if not math.isclose(repaired['RQ2']['F13'].value,p,rel_tol=1e-10):raise ValueError('ANOVA p value differs.')
    for name in ['Q1','Q2','Q3','Q4','Q5','RQ1']:
        old=list(original[name].values);new=list(repaired[name].values)
        if len(old)!=len(new):raise ValueError(f'Original sheet dimensions changed: {name}')
        for row_a,row_b in zip(old,new):
            if len(row_a)!=len(row_b):raise ValueError(f'Original sheet dimensions changed: {name}')
            for a,b in zip(row_a,row_b):
                equivalent=math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-9) if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==b
                if not equivalent:raise ValueError(f'Original finding changed in {name}: {a!r} -> {b!r}')
        if len(original[name]._pivots)!=len(repaired[name]._pivots):raise ValueError(f'Pivot definitions changed: {name}')
    return {'records':count,'groups':{g:len(v) for g,v in groups.items()},'churn_rate':sum(r[12] for r in rows)/count,'anova_F':f,'anova_p':p,'source_preserved':True}

if __name__=='__main__': print(json.dumps(verify(),indent=2))
